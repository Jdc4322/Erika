"""Lectura local de documentos adjuntos para Erika."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

MAX_TEXT_CHARS = 8000
SUPPORTED_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".pdf", ".txt", ".md", ".csv", ".docx"}


def extract_attachment(filename: str, data: bytes) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError("Formato no compatible. Usa imagen, PDF, DOCX, TXT, MD o CSV.")
    if len(data) > 15 * 1024 * 1024:
        raise ValueError("El archivo supera el l?mite de 15 MB.")

    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
        from PIL import Image
        try:
            with Image.open(BytesIO(data)) as image:
                image.verify()
        except Exception as exc:
            raise ValueError("No se pudo abrir esta imagen.") from exc
        return {"kind": "image", "name": Path(filename).name, "data": data}

    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        parts = []
        for page in reader.pages[:40]:
            parts.append(page.extract_text() or "")
        text = "\n\n".join(parts)
    elif suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8-sig", errors="replace")

    text = text.strip()
    if not text:
        raise ValueError("No encontr? texto legible en el archivo.")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS].rstrip() + "\n\n[Texto recortado por longitud.]"
    return {"kind": "document", "name": Path(filename).name, "text": text}
